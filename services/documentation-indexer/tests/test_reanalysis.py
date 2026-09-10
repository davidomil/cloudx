from __future__ import annotations

import io
import json
import sqlite3
import wave
from pathlib import Path

from fastapi.testclient import TestClient
import httpx
from openpyxl import Workbook
from PIL import Image
import pytest
from reportlab.pdfgen import canvas

import cloudx_documentation_indexer.archive as archive_module
from cloudx_documentation_indexer import DocumentationArchive, create_app
from cloudx_documentation_indexer.archive import ArchiveError
from cloudx_documentation_indexer.extraction import ExtractedSpan


def test_reanalysis_replaces_extraction_from_the_archived_source_without_duplicate_documents(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    source = tmp_path / "original.pdf"
    pdf = canvas.Canvas(str(source))
    pdf.drawString(40, 800, "Archived source contains ORIGINAL-EXTRACTION-19.")
    pdf.save()
    ingested = archive.ingest_path(source, title="Retained note", collection="board", tags=["reference"])[0]
    archive.invalidate_document(ingested.document_id, state="stale", reason="Retained review history.")
    archive.ingest_path(source, title="Retained note", collection="board", tags=["reference"])
    sibling = archive.ingest_upload(filename="sibling.pdf", content=source.read_bytes())
    source.unlink()
    archive.enrich_document(
        ingested.document_id,
        spans=[ExtractedSpan("Retained AI context contains PREVIOUS-ENRICHMENT-19.", "ai:metadata")],
        model="gpt-test",
        skill_ids=["documentation-enrich-metadata"],
    )
    before = archive.get_document(ingested.document_id)
    sibling_before = archive.get_document(sibling.document_id)
    archived_source = archive.root / before["snapshot_path"]
    original_bytes = archived_source.read_bytes()

    def updated_extractor(content, name, source_type, content_type, artifact_dir):
        assert content == original_bytes
        assert name == source.name
        assert source_type == "book"
        assert artifact_dir.parent != archived_source.parent
        artifact_dir.mkdir()
        (artifact_dir / "analysis.txt").write_text("New extraction artifact.", encoding="utf-8")
        return [ExtractedSpan("Updated extraction contains REANALYZED-CONTENT-19.", "page 1")]

    monkeypatch.setattr(archive_module, "extract_bytes", updated_extractor)
    for _ in range(2):
        result = archive.reanalyze_document(ingested.document_id)
        assert result.document_id == ingested.document_id
    after = archive.get_document(ingested.document_id)

    assert len(archive.list_documents()) == 2
    assert before["events"]
    assert len(list(archive.snapshots_dir.iterdir())) == 2
    for field in ["document_id", "title", "source_type", "uri", "content_sha256", "state", "collection", "tags_json", "created_at", "enrichments", "events"]:
        assert after[field] == before[field]
    assert archive.get_document(sibling.document_id) == sibling_before
    assert archived_source.read_bytes() == original_bytes
    assert (archive.root / after["snapshot_path"]).read_bytes() == original_bytes
    assert (archive.root / after["snapshot_path"]).parent.joinpath("extracted/analysis.txt").read_text() == "New extraction artifact."
    assert [chunk for chunk in after["chunks"] if chunk["chunk_origin"] == "ai"] == [chunk for chunk in before["chunks"] if chunk["chunk_origin"] == "ai"]
    assert {hit["documentId"] for hit in archive.search("REANALYZED-CONTENT-19", mode="lexical")} == {ingested.document_id}
    assert {hit["documentId"] for hit in archive.search("ORIGINAL-EXTRACTION-19", mode="lexical")} == {sibling.document_id}


def test_repeated_reanalysis_retains_only_the_current_unshared_snapshot(tmp_path: Path) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_text(text="Repeated analysis of RETAINED-SOURCE-19.")

    for _ in range(3):
        previous_snapshot = archive.root / archive.get_document(document.document_id)["snapshot_path"]
        archive.reanalyze_document(document.document_id)
        assert not previous_snapshot.exists()
        assert len(list(archive.snapshots_dir.iterdir())) == 1

    reopened = DocumentationArchive(archive.root)
    assert reopened.search("RETAINED-SOURCE-19", mode="lexical")[0]["documentId"] == document.document_id


def test_committed_reanalysis_survives_previous_snapshot_cleanup_failure(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_text(text="Published analysis retains PUBLISHED-SOURCE-19.")
    previous_snapshot = archive.get_document(document.document_id)["snapshot_path"]

    def fail_cleanup(_snapshot_path):
        raise OSError("forced snapshot cleanup failure")

    monkeypatch.setattr(archive, "_discard_unreferenced_snapshot", fail_cleanup)
    result = archive.reanalyze_document(document.document_id)

    assert result.document_id == document.document_id
    current = archive.get_document(document.document_id)
    assert current["snapshot_path"] != previous_snapshot
    assert (archive.root / current["snapshot_path"]).is_file()
    assert archive.search("PUBLISHED-SOURCE-19", mode="lexical")[0]["documentId"] == document.document_id
    assert "was published, but the previous snapshot could not be removed" in caplog.text


def test_reanalysis_reruns_pdf_extraction_after_the_original_file_is_removed(tmp_path: Path) -> None:
    source = tmp_path / "board.pdf"
    pdf = canvas.Canvas(str(source))
    pdf.drawString(40, 800, "PDF-REANALYSIS-19 uses a retained source snapshot.")
    pdf.save()
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_path(source, source_type="datasheet")[0]
    source.unlink()

    archive.reanalyze_document(document.document_id)

    record = archive.get_document(document.document_id)
    assert record["chunks"][0]["locator"] == "page 1"
    assert "PDF-REANALYSIS-19" in record["chunks"][0]["text"]
    assert archive.search("PDF-REANALYSIS-19", mode="lexical")[0]["documentId"] == document.document_id


def test_reanalysis_retains_committed_source_and_artifacts_when_projection_read_fails(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "board.pdf"
    pdf = canvas.Canvas(str(source))
    pdf.drawString(40, 800, "COMMITTED-SOURCE-19 survives projection read failure.")
    pdf.rect(40, 700, 120, 40)
    pdf.save()
    source_bytes = source.read_bytes()
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_path(source)[0]
    previous_snapshot = archive.root / archive.get_document(document.document_id)["snapshot_path"]
    artifacts = {
        artifact.relative_to(previous_snapshot.parent): artifact.read_bytes()
        for artifact in (previous_snapshot.parent / "extracted").rglob("*")
        if artifact.is_file()
    }
    assert artifacts
    source.unlink()

    def fail_projection_read():
        assert archive.get_document(document.document_id)["snapshot_path"] != previous_snapshot.relative_to(archive.root).as_posix()
        raise sqlite3.OperationalError("database is locked")

    with monkeypatch.context() as projection_failure:
        projection_failure.setattr(archive, "_projected_index_generation", fail_projection_read)
        with pytest.raises(sqlite3.OperationalError, match="database is locked"):
            archive.reanalyze_document(document.document_id)

    current = archive.get_document(document.document_id)
    committed_snapshot = archive.root / current["snapshot_path"]
    assert committed_snapshot != previous_snapshot
    assert committed_snapshot.read_bytes() == source_bytes
    for relative_path, content in artifacts.items():
        assert (committed_snapshot.parent / relative_path).read_bytes() == content
    assert archive.search("COMMITTED-SOURCE-19")[0]["documentId"] == document.document_id

    reopened = DocumentationArchive(archive.root)
    assert reopened.health()["ready"] is True
    assert reopened.reanalyze_document(document.document_id).document_id == document.document_id
    assert (archive.root / reopened.get_document(document.document_id)["snapshot_path"]).read_bytes() == source_bytes


@pytest.mark.parametrize(("uri", "source_type"), [
    ("https://example.com/board.png", "image"),
    ("https://example.com/board.xlsx", "spreadsheet"),
    ("https://example.com/board.html", "website"),
])
def test_reanalysis_extracts_copied_text_from_its_retained_format(tmp_path: Path, uri: str, source_type: str) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    text = "COPIED-SOURCE-19 retains the literal <board> marker."
    document = archive.ingest_text(text=text, uri=uri, collection="boards", tags=["copied"])
    archive.enrich_document(
        document.document_id,
        spans=[ExtractedSpan("Prior enrichment remains available.", "ai:metadata")],
        model="gpt-test",
        skill_ids=["documentation-enrich-metadata"],
    )
    before = archive.get_document(document.document_id)
    assert before["source_type"] == source_type

    for _ in range(2):
        result = archive.reanalyze_document(document.document_id)
        after = archive.get_document(document.document_id)
        assert result.document_id == document.document_id
        for field in ["document_id", "title", "source_type", "uri", "content_sha256", "collection", "tags_json"]:
            assert after[field] == before[field]
        assert (archive.root / after["snapshot_path"]).read_bytes() == text.encode("utf-8")
        assert [(chunk["locator"], chunk["text"]) for chunk in after["chunks"] if chunk["chunk_origin"] == "source"] == [("text", text)]
        assert [chunk for chunk in after["chunks"] if chunk["chunk_origin"] == "ai"] == [chunk for chunk in before["chunks"] if chunk["chunk_origin"] == "ai"]
        assert archive.search("COPIED-SOURCE-19", mode="lexical")[0]["documentId"] == document.document_id
    assert len(archive.list_documents()) == 1


def test_reanalysis_preserves_copied_text_that_begins_with_a_pdf_header(tmp_path: Path) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    text = "%PDF-1.7\nCOPIED-PDF-EXPLANATION-19 explains the PDF header as plain text."
    document = archive.ingest_text(text=text)
    before = archive.get_document(document.document_id)

    for _ in range(2):
        result = archive.reanalyze_document(document.document_id)
        after = archive.get_document(document.document_id)
        assert result.document_id == document.document_id
        for field in ["document_id", "title", "source_type", "uri", "content_sha256"]:
            assert after[field] == before[field]
        assert (archive.root / after["snapshot_path"]).read_bytes() == text.encode("utf-8")
        assert [(chunk["locator"], chunk["text"]) for chunk in after["chunks"]] == [("text", text)]
        assert archive.search("COPIED-PDF-EXPLANATION-19", mode="lexical")[0]["documentId"] == document.document_id
    assert len(archive.list_documents()) == 1


def test_reanalysis_preserves_copied_text_when_identical_html_shares_snapshot_metadata(tmp_path: Path) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    text = "COPIED-HTML-SIBLING-19 retains the literal <board> marker."
    document = archive.ingest_text(text=text, uri="https://example.com/board.html")
    sibling = archive.ingest_upload(filename="sibling.html", content=text.encode("utf-8"), content_type="text/html")
    before = archive.get_document(document.document_id)
    sibling_before = archive.get_document(sibling.document_id)
    snapshot = archive.root / before["snapshot_path"]
    assert snapshot.parent == (archive.root / sibling_before["snapshot_path"]).parent
    assert json.loads((snapshot.parent / "metadata.json").read_text())["contentType"] == "text/html"

    for _ in range(2):
        result = archive.reanalyze_document(document.document_id)
        after = archive.get_document(document.document_id)
        assert result.document_id == document.document_id
        for field in ["document_id", "title", "source_type", "uri", "content_sha256"]:
            assert after[field] == before[field]
        assert (archive.root / after["snapshot_path"]).read_bytes() == text.encode("utf-8")
        assert [(chunk["locator"], chunk["text"]) for chunk in after["chunks"]] == [("text", text)]
        assert archive.get_document(sibling.document_id) == sibling_before

    sibling_result = archive.reanalyze_document(sibling.document_id)
    sibling_after = archive.get_document(sibling.document_id)
    assert sibling_result.document_id == sibling.document_id
    assert [(chunk["locator"], chunk["text"]) for chunk in sibling_after["chunks"]] == [("html", "COPIED-HTML-SIBLING-19 retains the literal\nmarker.")]
    assert (archive.root / sibling_after["snapshot_path"]).read_bytes() == text.encode("utf-8")
    assert len(archive.list_documents()) == 2


@pytest.mark.parametrize("filename", ["note.md", "note", "note.json", "note.txt"])
@pytest.mark.parametrize("import_sibling", [False, True], ids=["upload-only", "shared-html-metadata"])
def test_reanalysis_preserves_plain_text_when_an_html_url_replaces_shared_metadata(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, filename: str, import_sibling: bool) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    text = "RETAINED-PLAIN-TEXT-19 keeps <board> and <script>EXAMPLE-SCRIPT-19</script> literally."
    source_bytes = text.encode("utf-8")
    upload = archive.ingest_upload(filename=filename, content=source_bytes, source_type="reference", content_type="text/plain")
    archive.enrich_document(
        upload.document_id,
        spans=[ExtractedSpan("Prior enrichment remains available.", "ai:metadata")],
        model="gpt-test",
        skill_ids=["documentation-enrich-metadata"],
    )
    before = archive.get_document(upload.document_id)
    snapshot = archive.root / before["snapshot_path"]
    assert [(chunk["locator"], chunk["text"]) for chunk in before["chunks"] if chunk["chunk_origin"] == "source"] == [("text", text)]
    if import_sibling:
        url = f"https://example.com/{filename}"
        response = httpx.Response(200, request=httpx.Request("GET", url), headers={"content-type": "text/html"}, content=source_bytes)
        monkeypatch.setattr(archive_module, "fetch_url_bytes", lambda _url, _limit: (response, source_bytes))
        sibling = archive.ingest_url(url)
        sibling_before = archive.get_document(sibling.document_id)
        assert sibling_before["snapshot_path"] == before["snapshot_path"]
        assert json.loads(snapshot.with_name("metadata.json").read_text())["contentType"] == "text/html"
        assert [(chunk["locator"], chunk["text"]) for chunk in sibling_before["chunks"]] == [("html", "RETAINED-PLAIN-TEXT-19 keeps\nand\nliterally.")]

    for _ in range(2):
        result = archive.reanalyze_document(upload.document_id)
        current = archive.get_document(upload.document_id)
        assert result.document_id == upload.document_id
        for field in ["document_id", "title", "uri", "source_type", "content_sha256", "state", "collection", "tags_json", "created_at", "enrichments", "events"]:
            assert current[field] == before[field]
        assert (archive.root / current["snapshot_path"]).read_bytes() == source_bytes
        assert [(chunk["locator"], chunk["text"]) for chunk in current["chunks"] if chunk["chunk_origin"] == "source"] == [("text", text)]
        assert [chunk for chunk in current["chunks"] if chunk["chunk_origin"] == "ai"] == [chunk for chunk in before["chunks"] if chunk["chunk_origin"] == "ai"]
        assert {hit["documentId"] for hit in archive.search("EXAMPLE-SCRIPT-19", mode="lexical")} == {upload.document_id}
        if import_sibling:
            assert archive.get_document(sibling.document_id) == sibling_before
            assert snapshot.read_bytes() == source_bytes
    assert len(archive.list_documents()) == (2 if import_sibling else 1)


@pytest.mark.parametrize("filename", ["scan", "scan.bin", "scan.txt"])
def test_reanalysis_preserves_explicit_image_extraction_for_unrecognized_filenames(tmp_path: Path, filename: str) -> None:
    source = tmp_path / filename
    Image.new("RGB", (20, 10), "white").save(source, format="PNG")
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_path(source, source_type="image")[0]
    before = archive.get_document(document.document_id)
    source_bytes = source.read_bytes()
    source.unlink()

    result = archive.reanalyze_document(document.document_id)

    after = archive.get_document(document.document_id)
    assert result.document_id == document.document_id
    assert after["source_type"] == "image"
    assert [(chunk["locator"], chunk["text"]) for chunk in after["chunks"]] == [(chunk["locator"], chunk["text"]) for chunk in before["chunks"]]
    snapshot = archive.root / after["snapshot_path"]
    assert snapshot.read_bytes() == source_bytes
    metadata = json.loads((snapshot.parent / "extracted/image_metadata.json").read_text(encoding="utf-8"))
    assert (metadata["format"], metadata["width"], metadata["height"]) == ("PNG", 20, 10)
    assert (snapshot.parent / "extracted" / metadata["artifact"]).is_file()


def test_reanalysis_preserves_explicit_html_extraction_for_a_txt_original(tmp_path: Path) -> None:
    source = tmp_path / "page.txt"
    source.write_text("<h1>ORIGINAL-WEBSITE-19</h1><script>HIDDEN-SCRIPT-19</script>", encoding="utf-8")
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_path(source, source_type="website")[0]
    before = archive.get_document(document.document_id)
    source.unlink()

    archive.reanalyze_document(document.document_id)

    after = archive.get_document(document.document_id)
    assert after["document_id"] == before["document_id"]
    assert after["source_type"] == "website"
    assert [(chunk["locator"], chunk["text"]) for chunk in after["chunks"]] == [("html", "ORIGINAL-WEBSITE-19")]


@pytest.mark.parametrize("content", [
    '["LOCAL-METADATA-SOURCE-19", 1]',
    '{"youtube": "LOCAL-METADATA-SOURCE-19"}',
])
def test_reanalysis_does_not_interpret_a_source_named_metadata_json_as_a_sidecar(tmp_path: Path, content: str) -> None:
    source = tmp_path / "metadata.json"
    source.write_text(content, encoding="utf-8")
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_path(source)[0]
    before = archive.get_document(document.document_id)
    source.unlink()

    for _ in range(2):
        result = archive.reanalyze_document(document.document_id)
        after = archive.get_document(document.document_id)
        assert result.document_id == document.document_id
        for field in ["document_id", "title", "source_type", "uri", "content_sha256"]:
            assert after[field] == before[field]
        assert Path(after["snapshot_path"]).name == "metadata.json"
        assert (archive.root / after["snapshot_path"]).read_bytes() == content.encode("utf-8")
        assert after["chunks"][0]["text"] == content
        assert archive.search("LOCAL-METADATA-SOURCE-19", mode="lexical")[0]["documentId"] == document.document_id
    assert len(archive.list_documents()) == 1


@pytest.mark.parametrize("content", [
    '["SHARED-METADATA-SOURCE-19 retains <board> literally.", 1]',
    '{"contentType": 42, "youtube": "SHARED-METADATA-SOURCE-19"}',
    "SHARED-METADATA-SOURCE-19 contains plain text instead of JSON.",
])
@pytest.mark.parametrize("sibling_state", ["absent", "shared", "reanalyzed"])
def test_reanalysis_preserves_copied_text_when_a_sibling_source_is_named_metadata_json(tmp_path: Path, content: str, sibling_state: str) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_text(text=content, title="Copied JSON", collection="board", tags=["reference"])
    archive.enrich_document(
        document.document_id,
        spans=[ExtractedSpan("Retained prior enrichment.", "ai:metadata")],
        model="gpt-test",
        skill_ids=["documentation-enrich-metadata"],
    )
    before = archive.get_document(document.document_id)
    snapshot = archive.root / before["snapshot_path"]
    if sibling_state != "absent":
        source = tmp_path / "metadata.json"
        source.write_bytes(content.encode("utf-8"))
        sibling = archive.ingest_path(source)[0]
        source.unlink()
        assert (archive.root / archive.get_document(sibling.document_id)["snapshot_path"]).parent == snapshot.parent
        if sibling_state == "reanalyzed":
            archive.reanalyze_document(sibling.document_id)
        sibling_before = archive.get_document(sibling.document_id)
        assert snapshot.with_name("metadata.json").read_bytes() == content.encode("utf-8")

    for _ in range(2):
        archive = DocumentationArchive(archive.root)
        result = archive.reanalyze_document(document.document_id)
        after = archive.get_document(document.document_id)
        assert result.document_id == document.document_id
        for field in ["document_id", "title", "source_type", "uri", "content_sha256", "state", "collection", "tags_json", "created_at", "enrichments", "events"]:
            assert after[field] == before[field]
        current_snapshot = archive.root / after["snapshot_path"]
        assert current_snapshot.read_bytes() == content.encode("utf-8")
        assert not current_snapshot.with_name("metadata.json").exists()
        assert [(chunk["locator"], chunk["text"]) for chunk in after["chunks"] if chunk["chunk_origin"] == "source"] == [("text", content)]
        assert [chunk for chunk in after["chunks"] if chunk["chunk_origin"] == "ai"] == [chunk for chunk in before["chunks"] if chunk["chunk_origin"] == "ai"]
        expected_ids = {document.document_id}
        if sibling_state != "absent":
            assert archive.get_document(sibling.document_id) == sibling_before
            assert (archive.root / sibling_before["snapshot_path"]).read_bytes() == content.encode("utf-8")
            expected_ids.add(sibling.document_id)
        assert {hit["documentId"] for hit in archive.search("SHARED-METADATA-SOURCE-19", mode="lexical")} == expected_ids

    if sibling_state != "absent":
        for _ in range(2):
            archive = DocumentationArchive(archive.root)
            assert archive.reanalyze_document(sibling.document_id).document_id == sibling.document_id
            sibling_after = archive.get_document(sibling.document_id)
            assert (archive.root / sibling_after["snapshot_path"]).read_bytes() == content.encode("utf-8")
            assert [(chunk["locator"], chunk["text"]) for chunk in sibling_after["chunks"]] == [("text", content)]
            assert archive.get_document(document.document_id) == after
    assert len(archive.list_documents()) == (1 if sibling_state == "absent" else 2)


@pytest.mark.parametrize("filename", ["recording.bin", "recording"])
def test_reanalysis_preserves_media_upload_identity_when_a_url_replaces_shared_metadata(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, filename: str) -> None:
    audio = io.BytesIO()
    with wave.open(audio, "wb") as recording:
        recording.setnchannels(1)
        recording.setsampwidth(2)
        recording.setframerate(16000)
        recording.writeframes(b"\x00\x00" * 160)
    source_bytes = audio.getvalue()
    archive = DocumentationArchive(tmp_path / "archive")
    upload = archive.ingest_upload(filename=filename, content=source_bytes, source_type="media", content_type="application/octet-stream")
    before = archive.get_document(upload.document_id)
    shared_metadata = (archive.root / before["snapshot_path"]).parent / "metadata.json"
    assert json.loads(shared_metadata.read_text())["upload"] is True

    url = f"https://example.com/{filename}"
    response = httpx.Response(200, request=httpx.Request("GET", url), headers={"content-type": "application/octet-stream"}, content=source_bytes)
    monkeypatch.setattr(archive_module, "fetch_url_bytes", lambda _url, _limit: (response, source_bytes))
    sibling = archive.ingest_url(url)
    url_metadata = json.loads(shared_metadata.read_text())
    assert url_metadata["url"] == url
    assert "upload" not in url_metadata
    assert archive.get_document(sibling.document_id)["snapshot_path"] == before["snapshot_path"]

    for _ in range(2):
        result = archive.reanalyze_document(upload.document_id)
        current = archive.get_document(upload.document_id)
        snapshot = archive.root / current["snapshot_path"]
        assert result.document_id == upload.document_id
        for field in ["document_id", "uri", "source_type", "content_sha256"]:
            assert current[field] == before[field]
        assert current["uri"] == f"upload://{filename}"
        assert current["source_type"] == "media"
        assert snapshot.read_bytes() == source_bytes
        assert json.loads(snapshot.with_name("metadata.json").read_text()) == url_metadata
    assert len(archive.list_documents()) == 2


def test_reanalysis_preserves_upload_metadata_for_sources_without_filename_extensions(tmp_path: Path) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_upload(
        filename="download",
        content=b"<html><body><h1>HTML-REANALYSIS-19</h1><script>HIDDEN-SCRIPT-19</script></body></html>",
        content_type="text/html",
        source_type="reference",
    )
    before = archive.get_document(document.document_id)
    metadata_before = (archive.root / before["snapshot_path"]).parent.joinpath("metadata.json").read_bytes()

    archive.reanalyze_document(document.document_id)

    after = archive.get_document(document.document_id)
    assert after["chunks"][0]["locator"] == "html"
    assert "HTML-REANALYSIS-19" in after["chunks"][0]["text"]
    assert "HIDDEN-SCRIPT-19" not in after["chunks"][0]["text"]
    assert (archive.root / after["snapshot_path"]).parent.joinpath("metadata.json").read_bytes() == metadata_before


@pytest.mark.parametrize("source_type", ["book", "website"])
@pytest.mark.parametrize("import_sibling", [False, True], ids=["upload-only", "shared-url-metadata"])
def test_reanalysis_preserves_html_format_when_a_url_replaces_shared_metadata(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, source_type: str, import_sibling: bool) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    source_bytes = b"<html><body><h1>RETAINED-HTML-FORMAT-19</h1><script>HIDDEN-SCRIPT-19</script></body></html>"
    upload = archive.ingest_upload(filename="download", content=source_bytes, source_type=source_type, content_type="text/html")
    archive.enrich_document(
        upload.document_id,
        spans=[ExtractedSpan("Previous enrichment stays available.", "ai:metadata")],
        model="gpt-test",
        skill_ids=["documentation-enrich-metadata"],
    )
    before = archive.get_document(upload.document_id)
    snapshot = archive.root / before["snapshot_path"]
    assert [(chunk["locator"], chunk["text"]) for chunk in before["chunks"] if chunk["chunk_origin"] == "source"] == [("html", "RETAINED-HTML-FORMAT-19")]
    if import_sibling:
        url = "https://example.com/download"
        response = httpx.Response(200, request=httpx.Request("GET", url), headers={"content-type": "application/octet-stream"}, content=source_bytes)
        monkeypatch.setattr(archive_module, "fetch_url_bytes", lambda _url, _limit: (response, source_bytes))
        sibling = archive.ingest_url(url)
        sibling_before = archive.get_document(sibling.document_id)
        assert sibling_before["snapshot_path"] == before["snapshot_path"]
        assert json.loads(snapshot.with_name("metadata.json").read_text())["contentType"] == "application/octet-stream"

    for _ in range(2):
        result = archive.reanalyze_document(upload.document_id)
        current = archive.get_document(upload.document_id)
        assert result.document_id == upload.document_id
        for field in ["document_id", "title", "uri", "source_type", "content_sha256", "enrichments"]:
            assert current[field] == before[field]
        assert (archive.root / current["snapshot_path"]).read_bytes() == source_bytes
        source_chunks = [chunk for chunk in current["chunks"] if chunk["chunk_origin"] == "source"]
        assert [(chunk["locator"], chunk["text"]) for chunk in source_chunks] == [("html", "RETAINED-HTML-FORMAT-19")]
        assert all("HIDDEN-SCRIPT-19" not in chunk["text"] for chunk in source_chunks)
        assert [chunk for chunk in current["chunks"] if chunk["chunk_origin"] == "ai"] == [chunk for chunk in before["chunks"] if chunk["chunk_origin"] == "ai"]
        if import_sibling:
            assert archive.get_document(sibling.document_id) == sibling_before
            assert snapshot.read_bytes() == source_bytes
    assert len(archive.list_documents()) == (2 if import_sibling else 1)


@pytest.mark.parametrize("source_format", ["png", "xlsx", "xls", "ods"])
@pytest.mark.parametrize("import_sibling", [False, True], ids=["upload-only", "shared-url-metadata"])
def test_reanalysis_preserves_mime_selected_formats_when_a_url_replaces_shared_metadata(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, source_format: str, import_sibling: bool) -> None:
    source = io.BytesIO()
    if source_format == "png":
        Image.new("RGB", (20, 10), "red").save(source, format="PNG")
        content_type = "image/png"
        expected_locators = {"image"}
        expected_artifacts = {"image_metadata.json", "images/download.png"}
    else:
        rows = [["Component", "Current"], ["MCU", 45], ["Sensor", 12]]
        if source_format == "xlsx":
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Power Budget"
            for row in rows:
                sheet.append(row)
            sheet.append(["Total", "=SUM(B2:B3)"])
            sheet["A6"] = "Merged board notes"
            sheet.merge_cells("A6:B6")
            workbook.save(source)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        elif source_format == "xls":
            import xlwt

            workbook = xlwt.Workbook()
            sheet = workbook.add_sheet("Power Budget")
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    sheet.write(row_index, column_index, value)
            workbook.save(source)
            content_type = "application/vnd.ms-excel"
        else:
            import pandas as pd

            pd.DataFrame(rows).to_excel(source, engine="odf", sheet_name="Power Budget", header=False, index=False)
            content_type = "application/vnd.oasis.opendocument.spreadsheet"
        expected_locators = {f"sheet Power Budget range A1:B{6 if source_format == 'xlsx' else 3}"}
        expected_artifacts = {
            "spreadsheet_index.tsv",
            "spreadsheets/sheet-001-Power_Budget.csv",
            "spreadsheets/sheet-001-Power_Budget.md",
            "spreadsheets/sheet-001-Power_Budget.json",
        }
    source_bytes = source.getvalue()
    archive = DocumentationArchive(tmp_path / "archive")
    filename = "download" if source_format in {"png", "xlsx"} else f"download.{source_format}"
    upload = archive.ingest_upload(filename=filename, content=source_bytes, source_type="reference", content_type=content_type)
    archive.enrich_document(
        upload.document_id,
        spans=[ExtractedSpan("Previous enrichment stays available.", "ai:metadata")],
        model="gpt-test",
        skill_ids=["documentation-enrich-metadata"],
    )
    before = archive.get_document(upload.document_id)
    snapshot = archive.root / before["snapshot_path"]
    source_chunks = [(chunk["locator"], chunk["text"]) for chunk in before["chunks"] if chunk["chunk_origin"] == "source"]
    assert {locator for locator, _ in source_chunks} == expected_locators
    artifacts = {
        artifact.relative_to(snapshot.parent / "extracted").as_posix(): artifact.read_bytes()
        for artifact in (snapshot.parent / "extracted").rglob("*") if artifact.is_file()
    }
    assert set(artifacts) == expected_artifacts
    if source_format == "xlsx":
        table = json.loads(artifacts["spreadsheets/sheet-001-Power_Budget.json"])
        assert table["formulas"] == [{"cell": "B4", "formula": "=SUM(B2:B3)"}]
        assert table["mergedRanges"] == ["A6:B6"]
    if import_sibling:
        url = f"https://example.com/{filename}"
        response = httpx.Response(200, request=httpx.Request("GET", url), headers={"content-type": "application/octet-stream"}, content=source_bytes)
        monkeypatch.setattr(archive_module, "fetch_url_bytes", lambda _url, _limit: (response, source_bytes))
        sibling = archive.ingest_url(url)
        sibling_before = archive.get_document(sibling.document_id)
        assert sibling_before["snapshot_path"] == before["snapshot_path"]
        assert json.loads(snapshot.with_name("metadata.json").read_text())["contentType"] == "application/octet-stream"

    for _ in range(2):
        previous_snapshot = archive.root / archive.get_document(upload.document_id)["snapshot_path"]
        result = archive.reanalyze_document(upload.document_id)
        current = archive.get_document(upload.document_id)
        replacement_snapshot = archive.root / current["snapshot_path"]
        assert result.document_id == upload.document_id
        for field in ["document_id", "title", "uri", "source_type", "content_sha256", "enrichments"]:
            assert current[field] == before[field]
        assert replacement_snapshot != previous_snapshot
        assert replacement_snapshot.read_bytes() == source_bytes
        assert [(chunk["locator"], chunk["text"]) for chunk in current["chunks"] if chunk["chunk_origin"] == "source"] == source_chunks
        assert [chunk for chunk in current["chunks"] if chunk["chunk_origin"] == "ai"] == [chunk for chunk in before["chunks"] if chunk["chunk_origin"] == "ai"]
        regenerated_artifacts = {
            artifact.relative_to(replacement_snapshot.parent / "extracted").as_posix(): artifact.read_bytes()
            for artifact in (replacement_snapshot.parent / "extracted").rglob("*") if artifact.is_file()
        }
        assert regenerated_artifacts == artifacts
        if import_sibling:
            assert archive.get_document(sibling.document_id) == sibling_before
            assert snapshot.read_bytes() == source_bytes
    assert len(archive.list_documents()) == (2 if import_sibling else 1)


@pytest.mark.parametrize("failure", ["extraction", "empty extraction", "index publication"])
def test_failed_reanalysis_preserves_the_complete_previous_archive(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, failure: str) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    source = io.BytesIO()
    pdf = canvas.Canvas(source)
    pdf.drawString(40, 800, "Preserved source contains PRESERVED-SOURCE-19.")
    pdf.save()
    document = archive.ingest_upload(filename="preserved.pdf", title="Preserved analysis", content=source.getvalue())
    archive.enrich_document(
        document.document_id,
        spans=[ExtractedSpan("Previous enrichment contains PRESERVED-AI-19.", "ai:metadata")],
        model="gpt-test",
        skill_ids=["documentation-enrich-metadata"],
    )
    before = archive.get_document(document.document_id)
    published_files = {entry["path"]: entry["sha256"] for entry in archive.portable_manifest()["files"]}

    def extract(content, name, source_type, content_type, artifact_dir):
        artifact_dir.mkdir()
        (artifact_dir / "incomplete.txt").write_text("Do not publish.", encoding="utf-8")
        if failure == "extraction":
            raise RuntimeError("forced extraction failure")
        if failure == "empty extraction":
            return []
        return [ExtractedSpan("Unpublished replacement UNPUBLISHED-REANALYSIS-19.", "page 1")]

    def fail_index_write(_index, _path):
        raise RuntimeError("forced index failure")

    monkeypatch.setattr(archive_module, "extract_bytes", extract)
    if failure == "index publication":
        monkeypatch.setattr(archive_module.IdMapIndex, "write", fail_index_write)
    with pytest.raises((ArchiveError, RuntimeError)):
        archive.reanalyze_document(document.document_id)

    assert archive.get_document(document.document_id) == before
    assert {entry["path"]: entry["sha256"] for entry in archive.portable_manifest()["files"]} == published_files
    assert archive.search("PRESERVED-SOURCE-19", mode="lexical")[0]["documentId"] == document.document_id
    assert archive.search("PRESERVED-AI-19", mode="lexical")[0]["documentId"] == document.document_id
    assert archive.search("UNPUBLISHED-REANALYSIS-19", mode="lexical") == []


@pytest.mark.parametrize("state", ["stale", "deleted", "superseded"])
def test_reanalysis_does_not_reactivate_inactive_documents(tmp_path: Path, state: str) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_text(text="Inactive source.")
    archive.invalidate_document(document.document_id, state=state, reason="Keep excluded.")
    before = archive.get_document(document.document_id)

    with pytest.raises(ArchiveError, match="Only active documents"):
        archive.reanalyze_document(document.document_id)

    assert archive.get_document(document.document_id) == before


@pytest.mark.parametrize("problem", ["traversal", "absolute", "symlink", "missing", "content hash", "metadata escape", "invalid metadata", "invalid metadata JSON", "invalid metadata encoding", "invalid content type"])
def test_reanalysis_rejects_invalid_archived_sources_before_extraction(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, problem: str) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_text(text="Archived source.")
    snapshot = archive.root / archive.get_document(document.document_id)["snapshot_path"]
    outside = tmp_path / "outside.txt"
    outside.write_text("Outside source.", encoding="utf-8")
    if problem in {"traversal", "absolute"}:
        with archive._connect() as db:
            db.execute("UPDATE documents SET snapshot_path = ? WHERE document_id = ?", ("../outside.txt" if problem == "traversal" else str(outside), document.document_id))
    elif problem == "symlink":
        snapshot.unlink()
        snapshot.symlink_to(outside)
    elif problem == "missing":
        snapshot.unlink()
    elif problem == "content hash":
        snapshot.write_text("Changed source.", encoding="utf-8")
    elif problem == "metadata escape":
        (snapshot.parent / "metadata.json").symlink_to(outside)
    elif problem == "invalid content type":
        (snapshot.parent / "metadata.json").write_text('{"contentType": 42}', encoding="utf-8")
    elif problem == "invalid metadata JSON":
        (snapshot.parent / "metadata.json").write_text("invalid JSON", encoding="utf-8")
    elif problem == "invalid metadata encoding":
        (snapshot.parent / "metadata.json").write_bytes(b'\xff')
    else:
        (snapshot.parent / "metadata.json").write_text("[]", encoding="utf-8")

    def unexpected_extraction(*args):
        pytest.fail("Invalid source reached extraction.")

    monkeypatch.setattr(archive_module, "extract_bytes", unexpected_extraction)
    with pytest.raises(ArchiveError):
        archive.reanalyze_document(document.document_id)
    assert not list(archive.snapshots_dir.glob("reanalysis-*"))


@pytest.fixture
def youtube_archive(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[DocumentationArchive, archive_module.IngestedDocument]:
    metadata = archive_module.YouTubeVideoMetadata(
        title="Retained lecture",
        webpage_url="https://www.youtube.com/watch?v=retained-lecture",
        stream_url="mock://retained-lecture",
        http_headers={},
        duration=60,
    )

    def acquire_evidence(_url, _metadata, artifact_dir, *, progress=None):
        frames_dir = artifact_dir / "media" / "keyframes"
        frames_dir.mkdir(parents=True)
        Image.new("RGB", (64, 36), "white").save(frames_dir / "frame-000001.jpg")
        return [archive_module.TranscriptSegment(12.0, 18.0, "SHARED-YOUTUBE-COPY-19 retains <strong>timed transcript evidence.</strong>")], [{
            "offsetSeconds": 12,
            "path": "media/keyframes/frame-000001.jpg",
            "reason": "segment-start",
            "transcriptStartSeconds": 12.0,
            "transcriptEndSeconds": 18.0,
        }]

    monkeypatch.setattr(archive_module, "extract_youtube_video_metadata", lambda _url: metadata)
    monkeypatch.setattr(archive_module, "extract_youtube_video_evidence", acquire_evidence)
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_youtube_video(metadata.webpage_url)
    archive.enrich_document(
        document.document_id,
        spans=[ExtractedSpan("Prior video enrichment remains available.", "ai:media")],
        model="gpt-test",
        skill_ids=["documentation-enrich-media"],
    )
    return archive, document


def test_generated_code_requires_ai_reenrichment_without_replacing_its_structure(tmp_path: Path) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    document = archive.ingest_upload(filename="driver.c", content=b"void reset(void) {}", accept_generated_code_documentation=True)
    before = archive.get_document(document.document_id)

    with pytest.raises(ArchiveError, match="Rerun AI enrichment"):
        archive.reanalyze_document(document.document_id)

    assert archive.get_document(document.document_id) == before


def test_reanalysis_preserves_copied_text_when_generated_youtube_shares_snapshot_metadata(youtube_archive) -> None:
    archive, generated = youtube_archive
    generated_before = archive.get_document(generated.document_id)
    generated_snapshot = archive.root / generated_before["snapshot_path"]
    source_bytes = generated_snapshot.read_bytes()
    copied = archive.ingest_text(title="Manual lecture reference", text=source_bytes.decode("utf-8"))
    archive.enrich_document(
        copied.document_id,
        spans=[ExtractedSpan("Prior copied-text enrichment remains available.", "media metadata")],
        model="gpt-test",
        skill_ids=["documentation-enrich-media"],
    )
    before = archive.get_document(copied.document_id)
    source_chunks = [(chunk["locator"], chunk["text"]) for chunk in before["chunks"] if chunk["chunk_origin"] == "source"]
    assert {locator for locator, _ in source_chunks} == {"text"}
    snapshot = archive.root / before["snapshot_path"]
    assert snapshot.parent == generated_snapshot.parent
    assert "youtube" in json.loads((snapshot.parent / "metadata.json").read_text())
    assert {chunk["locator"] for chunk in generated_before["chunks"] if chunk["chunk_origin"] == "source"} == {
        "media metadata", "transcript 00:12-00:18", "media keyframe keyframe-000012 00:12",
    }
    generated_files = {path: path.read_bytes() for path in generated_snapshot.parent.rglob("*") if path.is_file()}

    for _ in range(2):
        result = archive.reanalyze_document(copied.document_id)
        archive = DocumentationArchive(archive.root)
        after = archive.get_document(copied.document_id)
        assert result.document_id == copied.document_id
        for field in ["document_id", "title", "source_type", "uri", "content_sha256", "created_at", "enrichments"]:
            assert after[field] == before[field]
        assert (archive.root / after["snapshot_path"]).read_bytes() == source_bytes
        assert [(chunk["locator"], chunk["text"]) for chunk in after["chunks"] if chunk["chunk_origin"] == "source"] == source_chunks
        assert [chunk for chunk in after["chunks"] if chunk["chunk_origin"] == "ai"] == [chunk for chunk in before["chunks"] if chunk["chunk_origin"] == "ai"]
        assert {hit["documentId"] for hit in archive.search("SHARED-YOUTUBE-COPY-19", mode="lexical")} == {copied.document_id, generated.document_id}
        assert archive.get_document(generated.document_id) == generated_before
        assert {path: path.read_bytes() for path in generated_files} == generated_files
        assert len(archive.list_documents()) == 2


@pytest.mark.parametrize("shared_archive", [False, True], ids=["separate-archive", "shared-youtube-metadata"])
def test_reanalysis_preserves_local_html_when_generated_youtube_shares_snapshot_metadata(tmp_path: Path, youtube_archive, shared_archive: bool) -> None:
    video_archive, generated = youtube_archive
    generated_before = video_archive.get_document(generated.document_id)
    generated_snapshot = video_archive.root / generated_before["snapshot_path"]
    source_bytes = generated_snapshot.read_bytes()
    source = tmp_path / "lecture.html"
    source.write_bytes(source_bytes)
    archive = video_archive if shared_archive else DocumentationArchive(tmp_path / "html-archive")
    document = archive.ingest_path(source, collection="lectures", tags=["reference"])[0]
    source.unlink()
    archive.enrich_document(
        document.document_id,
        spans=[ExtractedSpan("Prior HTML enrichment remains available.", "media metadata")],
        model="gpt-test",
        skill_ids=["documentation-enrich-metadata"],
    )
    before = archive.get_document(document.document_id)
    source_chunks = [(chunk["locator"], chunk["text"]) for chunk in before["chunks"] if chunk["chunk_origin"] == "source"]
    assert before["source_type"] == "website"
    assert {locator for locator, _ in source_chunks} == {"html"}
    assert "<strong>" in source_bytes.decode("utf-8")
    assert all("<strong>" not in text for _, text in source_chunks)
    snapshot = archive.root / before["snapshot_path"]
    if shared_archive:
        assert snapshot.parent == generated_snapshot.parent
        assert "youtube" in json.loads(snapshot.with_name("metadata.json").read_text())
    else:
        assert not snapshot.with_name("metadata.json").exists()
    assert {chunk["locator"] for chunk in generated_before["chunks"] if chunk["chunk_origin"] == "source"} == {
        "media metadata", "transcript 00:12-00:18", "media keyframe keyframe-000012 00:12",
    }
    keyframe = generated_snapshot.parent / "extracted/media/keyframes/frame-000001.jpg"
    with Image.open(keyframe) as frame:
        assert frame.format == "JPEG"
        frame.verify()
    generated_files = {path: path.read_bytes() for path in generated_snapshot.parent.rglob("*") if path.is_file()}

    for _ in range(2):
        archive = DocumentationArchive(archive.root)
        result = archive.reanalyze_document(document.document_id)
        archive = DocumentationArchive(archive.root)
        after = archive.get_document(document.document_id)
        assert result.document_id == document.document_id
        for field in ["document_id", "title", "source_type", "uri", "content_sha256", "state", "collection", "tags_json", "created_at", "enrichments", "events"]:
            assert after[field] == before[field]
        assert (archive.root / after["snapshot_path"]).read_bytes() == source_bytes
        assert [(chunk["locator"], chunk["text"]) for chunk in after["chunks"] if chunk["chunk_origin"] == "source"] == source_chunks
        assert [chunk for chunk in after["chunks"] if chunk["chunk_origin"] == "ai"] == [chunk for chunk in before["chunks"] if chunk["chunk_origin"] == "ai"]
        expected_ids = {document.document_id, generated.document_id} if shared_archive else {document.document_id}
        assert {hit["documentId"] for hit in archive.search("SHARED-YOUTUBE-COPY-19", mode="lexical")} == expected_ids
        assert len(archive.list_documents()) == len(expected_ids)
        video_archive = DocumentationArchive(video_archive.root)
        assert video_archive.get_document(generated.document_id) == generated_before
        assert {path: path.read_bytes() for path in generated_files} == generated_files


def test_reanalysis_allows_copied_youtube_transcripts_with_media_locators_only_in_ai_chunks(tmp_path: Path) -> None:
    archive = DocumentationArchive(tmp_path / "archive")
    transcript = "Copied video transcript retains COPIED-YOUTUBE-19."
    document = archive.ingest_url(
        "https://www.youtube.com/watch?v=copied-transcript",
        title="Lecture.youtube",
        transcript=transcript,
    )
    archive.enrich_document(
        document.document_id,
        spans=[ExtractedSpan("An AI locator is not source provenance.", "media metadata")],
        model="gpt-test",
        skill_ids=["documentation-enrich-media"],
    )
    before = archive.get_document(document.document_id)

    for _ in range(2):
        assert archive.reanalyze_document(document.document_id).document_id == document.document_id
        after = archive.get_document(document.document_id)
        assert (archive.root / after["snapshot_path"]).read_text() == transcript
        assert [(chunk["locator"], chunk["text"]) for chunk in after["chunks"] if chunk["chunk_origin"] == "source"] == [("text", transcript)]
        assert [chunk for chunk in after["chunks"] if chunk["chunk_origin"] == "ai"] == [chunk for chunk in before["chunks"] if chunk["chunk_origin"] == "ai"]


def test_reanalysis_endpoint_returns_existing_document_identity_and_validates_state(tmp_path: Path) -> None:
    client = TestClient(create_app(tmp_path / "archive"))
    document = client.post("/ingest/text", json={"text": "API-REANALYSIS-19 retained transcript.", "sourceType": "media"}).json()["document"]
    document_id = document["documentId"]

    response = client.post(f"/documents/{document_id}/reanalyze")

    assert response.status_code == 200
    assert response.json() == {"documents": [document]}
    assert client.post("/documents/unknown/reanalyze").status_code == 400
    client.delete(f"/documents/{document_id}")
    response = client.post(f"/documents/{document_id}/reanalyze")
    assert response.status_code == 400
    assert "Only active documents" in response.json()["detail"]
